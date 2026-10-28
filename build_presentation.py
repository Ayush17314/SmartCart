import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_smartcart_presentation(output_path):
    prs = Presentation()
    # 16:9 Widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_slide_layout = prs.slide_layouts[6] # completely blank layout

    # Color Palette
    NAVY = RGBColor(15, 23, 42)        # #0F172A - Deep Slate
    BLUE = RGBColor(37, 99, 235)       # #2563EB - Royal Blue
    ACCENT_BLUE = RGBColor(30, 58, 138)# #1E3A8A - Navy Accent
    DARK = RGBColor(30, 41, 59)        # #1E293B - Text Dark
    MUTED = RGBColor(100, 116, 139)    # #64748B - Text Muted
    LIGHT_BG = RGBColor(248, 250, 252) # #F8FAFC - Soft card bg
    CARD_BORDER = RGBColor(226, 232, 240) # #E2E8F0 - Border line
    WHITE = RGBColor(255, 255, 255)
    GREEN = RGBColor(22, 163, 74)      # #16A34A - Success Green

    def add_header(slide, title_text, category_text=None):
        # Top banner/category
        if category_text:
            tb_cat = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.35))
            tf_cat = tb_cat.text_frame
            tf_cat.word_wrap = True
            p_cat = tf_cat.paragraphs[0]
            p_cat.text = category_text.upper()
            p_cat.font.name = "Calibri"
            p_cat.font.size = Pt(10)
            p_cat.font.bold = True
            p_cat.font.color.rgb = BLUE
            
        # Title
        tb_t = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.7), Inches(0.8))
        tf_t = tb_t.text_frame
        tf_t.word_wrap = True
        p_t = tf_t.paragraphs[0]
        p_t.text = title_text
        p_t.font.name = "Calibri"
        p_t.font.size = Pt(24)
        p_t.font.bold = True
        p_t.font.color.rgb = NAVY

    def create_card(slide, left, top, width, height, bg_color=LIGHT_BG, border_color=CARD_BORDER):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        if border_color:
            shape.line.color.rgb = border_color
            shape.line.width = Pt(1.5)
        else:
            shape.line.fill.background()
        return shape

    # =========================================================================
    # SLIDE 1: TITLE SLIDE
    # =========================================================================
    slide1 = prs.slides.add_slide(blank_slide_layout)
    
    # Background card decoration
    bg1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = WHITE
    bg1.line.fill.background()

    # Institute name at top
    tb_inst = slide1.shapes.add_textbox(Inches(1.0), Inches(0.6), Inches(11.333), Inches(0.6))
    p_inst = tb_inst.text_frame.paragraphs[0]
    p_inst.text = "INDIAN INSTITUTE OF INFORMATION TECHNOLOGY VADODARA"
    p_inst.font.name = "Calibri"
    p_inst.font.size = Pt(15)
    p_inst.font.bold = True
    p_inst.font.color.rgb = ACCENT_BLUE
    p_inst.alignment = PP_ALIGN.CENTER

    # Sub-heading
    tb_sub = slide1.shapes.add_textbox(Inches(1.0), Inches(1.2), Inches(11.333), Inches(0.4))
    p_sub = tb_sub.text_frame.paragraphs[0]
    p_sub.text = "Project Presentation on"
    p_sub.font.name = "Calibri"
    p_sub.font.size = Pt(13)
    p_sub.font.color.rgb = MUTED
    p_sub.alignment = PP_ALIGN.CENTER

    # Project Title
    tb_title = slide1.shapes.add_textbox(Inches(1.0), Inches(1.6), Inches(11.333), Inches(1.1))
    p_title = tb_title.text_frame.paragraphs[0]
    p_title.text = "SmartCart: Distributed E-Commerce Platform"
    p_title.font.name = "Calibri"
    p_title.font.size = Pt(32)
    p_title.font.bold = True
    p_title.font.color.rgb = NAVY
    p_title.alignment = PP_ALIGN.CENTER

    tb_desc = slide1.shapes.add_textbox(Inches(1.0), Inches(2.65), Inches(11.333), Inches(0.5))
    p_desc = tb_desc.text_frame.paragraphs[0]
    p_desc.text = "A Polyglot Microservices Web Application with Razorpay Payment Gateway Integration"
    p_desc.font.name = "Calibri"
    p_desc.font.size = Pt(13)
    p_desc.font.italic = True
    p_desc.font.color.rgb = BLUE
    p_desc.alignment = PP_ALIGN.CENTER

    # Center metadata container
    card_meta = create_card(slide1, Inches(2.2), Inches(3.4), Inches(8.933), Inches(2.6), bg_color=LIGHT_BG)
    tb_m = slide1.shapes.add_textbox(Inches(2.5), Inches(3.6), Inches(8.333), Inches(2.2))
    tf_m = tb_m.text_frame
    tf_m.word_wrap = True

    p1 = tf_m.paragraphs[0]
    p1.text = "Presented by:"
    p1.font.bold = True
    p1.font.size = Pt(12)
    p1.font.color.rgb = NAVY

    p2 = tf_m.add_paragraph()
    p2.text = "Ayush Dubey (Group 03 Lead, TeamID 03, Enrollment ID: 202351000)"
    p2.font.size = Pt(13)
    p2.font.color.rgb = DARK
    p2.space_after = Pt(12)

    p3 = tf_m.add_paragraph()
    p3.text = "Guide: Course Faculty & Lab Coordinator"
    p3.font.bold = True
    p3.font.size = Pt(12)
    p3.font.color.rgb = NAVY

    p4 = tf_m.add_paragraph()
    p4.text = "Department: Department of Computer Science & Engineering"
    p4.font.size = Pt(12)
    p4.font.color.rgb = DARK
    p4.space_after = Pt(12)

    p5 = tf_m.add_paragraph()
    p5.text = "CS 545 Software Engineering  •  MCA Second Year  •  Autumn 2026-27"
    p5.font.bold = True
    p5.font.size = Pt(11)
    p5.font.color.rgb = BLUE

    # =========================================================================
    # SLIDE 2: PROBLEM AND OBJECTIVES
    # =========================================================================
    slide2 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide2, "Problem and objectives", "SmartCart / Project Scope")

    # Left Column: Current Problem
    card_prob = create_card(slide2, Inches(0.8), Inches(1.7), Inches(5.6), Inches(5.1))
    tb_prob = slide2.shapes.add_textbox(Inches(1.1), Inches(1.9), Inches(5.0), Inches(4.7))
    tf_prob = tb_prob.text_frame
    tf_prob.word_wrap = True
    
    p_ph = tf_prob.paragraphs[0]
    p_ph.text = "CURRENT PROBLEM"
    p_ph.font.name = "Calibri"
    p_ph.font.size = Pt(14)
    p_ph.font.bold = True
    p_ph.font.color.rgb = RGBColor(220, 38, 38) # Red Accent
    p_ph.space_after = Pt(14)

    prob_items = [
        ("Single Point of Failure in Monoliths:", "Traditional e-commerce backends couple cart, catalog, and billing into a single database; high traffic during flash sales creates severe bottlenecks and system-wide outages."),
        ("Vulnerabilities in Cart & Payments:", "Client-side price calculations allow cart tampering, while unverified third-party payment callbacks cause fraudulent orders and inventory loss."),
        ("Slow, Error-Prone Order Operations:", "Tracking orders across disparate legacy modules, uncoordinated stock deductions, and manual return evaluations lead to stock discrepancies and customer dissatisfaction.")
    ]
    for title, desc in prob_items:
        p_t = tf_prob.add_paragraph()
        p_t.text = "• " + title + " "
        p_t.font.name = "Calibri"
        p_t.font.bold = True
        p_t.font.size = Pt(11.5)
        p_t.font.color.rgb = NAVY
        
        run_d = p_t.add_run()
        run_d.text = desc
        run_d.font.bold = False
        run_d.font.size = Pt(11)
        run_d.font.color.rgb = DARK
        p_t.space_after = Pt(12)

    # Right Column: Project Objectives
    card_obj = create_card(slide2, Inches(6.9), Inches(1.7), Inches(5.6), Inches(5.1))
    tb_obj = slide2.shapes.add_textbox(Inches(7.2), Inches(1.9), Inches(5.0), Inches(4.7))
    tf_obj = tb_obj.text_frame
    tf_obj.word_wrap = True
    
    p_oh = tf_obj.paragraphs[0]
    p_oh.text = "PROJECT OBJECTIVES"
    p_oh.font.name = "Calibri"
    p_oh.font.size = Pt(14)
    p_oh.font.bold = True
    p_oh.font.color.rgb = GREEN
    p_oh.space_after = Pt(14)

    obj_items = [
        ("Polyglot Microservices Architecture:", "Deconstruct e-commerce capabilities into 8 autonomous microservices (Auth, Product, Cart, Wishlist, Order, Payment, Notification, API Gateway) with isolated databases."),
        ("Atomic Checkout Saga with Rollbacks:", "Implement an atomic checkout workflow (placeOrder) that reserves warehouse stock, verifies coupons, and triggers compensatory stock rollbacks if payments fail."),
        ("Cryptographic Payments & Edge Resilience:", "Integrate Razorpay Orders (UPI/Cards in Paise) with HMAC-SHA256 signature verification, durable MySQL idempotency to block duplicate orders, and dynamic ETag caching.")
    ]
    for title, desc in obj_items:
        p_t = tf_obj.add_paragraph()
        p_t.text = "• " + title + " "
        p_t.font.name = "Calibri"
        p_t.font.bold = True
        p_t.font.size = Pt(11.5)
        p_t.font.color.rgb = NAVY
        
        run_d = p_t.add_run()
        run_d.text = desc
        run_d.font.bold = False
        run_d.font.size = Pt(11)
        run_d.font.color.rgb = DARK
        p_t.space_after = Pt(12)

    # =========================================================================
    # SLIDE 3: PROPOSED SOLUTION
    # =========================================================================
    slide3 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide3, "Proposed solution", "End-to-End Shopping & Fulfillment Flow")

    # Subtitle
    tb_sub3 = slide3.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.7), Inches(0.4))
    p_s3 = tb_sub3.text_frame.paragraphs[0]
    p_s3.text = "A customer discovers products, manages cart items, and completes verified checkout in a responsive web application."
    p_s3.font.name = "Calibri"
    p_s3.font.size = Pt(13)
    p_s3.font.color.rgb = DARK

    # 3 Step Cards
    step_data = [
        ("01", "Sign in & Discover", "Authenticate customer / staff", "JWT session tokens securely stored in HttpOnly SameSite=Lax cookies; browses catalog with full-text search, category filters, and dynamic 304 ETag caching."),
        ("02", "Manage Cart & Checkout", "Choose items and apply coupons", "Server-authoritative cart subtotal computation preventing client price tampering; validates stock availability and evaluates promotional discount codes."),
        ("03", "Pay & Track Order", "Validate payment & track shipment", "Multi-channel Razorpay checkout (UPI/Cards) in Indian Paise; validates HMAC-SHA256 signatures; advances order state (Confirmed -> Shipped -> Delivered).")
    ]

    card_w = Inches(3.64)
    gap = Inches(0.4)
    start_x = Inches(0.8)
    
    for i, (num, title, subtitle, desc) in enumerate(step_data):
        x = start_x + i * (card_w + gap)
        create_card(slide3, x, Inches(2.1), card_w, Inches(3.7))
        
        # Step Number Badge
        tb_step = slide3.shapes.add_textbox(x + Inches(0.25), Inches(2.3), Inches(3.1), Inches(3.3))
        tf_step = tb_step.text_frame
        tf_step.word_wrap = True
        
        p_n = tf_step.paragraphs[0]
        p_n.text = num
        p_n.font.name = "Calibri"
        p_n.font.size = Pt(28)
        p_n.font.bold = True
        p_n.font.color.rgb = BLUE
        p_n.space_after = Pt(4)
        
        p_tit = tf_step.add_paragraph()
        p_tit.text = title
        p_tit.font.name = "Calibri"
        p_tit.font.size = Pt(14)
        p_tit.font.bold = True
        p_tit.font.color.rgb = NAVY
        
        p_sub = tf_step.add_paragraph()
        p_sub.text = subtitle
        p_sub.font.name = "Calibri"
        p_sub.font.size = Pt(11)
        p_sub.font.bold = True
        p_sub.font.color.rgb = MUTED
        p_sub.space_after = Pt(10)
        
        p_desc = tf_step.add_paragraph()
        p_desc.text = desc
        p_desc.font.name = "Calibri"
        p_desc.font.size = Pt(11)
        p_desc.font.color.rgb = DARK

    # Output banner at bottom
    out_card = create_card(slide3, Inches(0.8), Inches(6.1), Inches(11.733), Inches(0.85), bg_color=WHITE, border_color=BLUE)
    tb_out = slide3.shapes.add_textbox(Inches(1.0), Inches(6.2), Inches(11.3), Inches(0.65))
    p_out = tb_out.text_frame.paragraphs[0]
    p_out.text = "Output: "
    p_out.font.name = "Calibri"
    p_out.font.size = Pt(12)
    p_out.font.bold = True
    p_out.font.color.rgb = BLUE
    
    r_out = p_out.add_run()
    r_out.text = "An itemized order confirmation receipt, transactional email notifications via Resend API, real-time shipment milestone tracking, and a 30-day return policy evaluation workflow."
    r_out.font.bold = False
    r_out.font.size = Pt(11.5)
    r_out.font.color.rgb = DARK

    # =========================================================================
    # SLIDE 4: SYSTEM ARCHITECTURE
    # =========================================================================
    slide4 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide4, "System architecture", "Polyglot Microservices Topology")

    # Subtitle
    tb_sub4 = slide4.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(11.7), Inches(0.35))
    p_s4 = tb_sub4.text_frame.paragraphs[0]
    p_s4.text = "A decoupled architecture that isolates business domains and scales horizontally"
    p_s4.font.name = "Calibri"
    p_s4.font.size = Pt(12.5)
    p_s4.font.color.rgb = DARK

    # Left: Three tier flow cards
    flow_steps = [
        ("Storefront / Client (Port 3000)", "Next.js 15, React 19, Tailwind CSS", "Customer storefront, cart drawer, checkout portal, and role-based staff admin panels with HttpOnly session cookies."),
        ("Central API Gateway (Port 4000)", "Reverse Proxy, Helmet & AJV Guard", "Sliding-window rate limiting (120 req/60s), durable MySQL idempotency key checks, and uniform RFC 7807 error standardization."),
        ("Microservices & DBs (Ports 4001–4007)", "Express.js & Independent MongoDB Stores", "Auth (4001), Product (4002), Cart (4003), Wishlist (4004), Order (4005), Payment (4006), and Notification (4007).")
    ]
    
    for i, (title, stack, desc) in enumerate(flow_steps):
        top_y = Inches(1.85 + i * 1.65)
        create_card(slide4, Inches(0.8), top_y, Inches(5.8), Inches(1.45))
        
        tb = slide4.shapes.add_textbox(Inches(1.0), top_y + Inches(0.1), Inches(5.4), Inches(1.25))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = title
        p.font.name = "Calibri"
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = NAVY
        
        p_st = tf.add_paragraph()
        p_st.text = stack
        p_st.font.name = "Calibri"
        p_st.font.size = Pt(10.5)
        p_st.font.bold = True
        p_st.font.color.rgb = BLUE
        p_st.space_after = Pt(2)
        
        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.name = "Calibri"
        p_d.font.size = Pt(10)
        p_d.font.color.rgb = DARK

    # Right: Embedded Architecture Diagram
    img_arch = os.path.join("docs", "Lab-1", "smartcart-architecture.png")
    if os.path.exists(img_arch):
        card_img = create_card(slide4, Inches(6.9), Inches(1.85), Inches(5.633), Inches(4.75), bg_color=WHITE)
        slide4.shapes.add_picture(img_arch, Inches(7.0), Inches(1.95), width=Inches(5.433))
        
        tb_cap = slide4.shapes.add_textbox(Inches(7.0), Inches(6.6), Inches(5.433), Inches(0.35))
        p_cap = tb_cap.text_frame.paragraphs[0]
        p_cap.text = "Request Journey: Gateway → Order Saga → Stock Hold → Razorpay → Resend Email"
        p_cap.font.name = "Calibri"
        p_cap.font.size = Pt(9.5)
        p_cap.font.bold = True
        p_cap.font.color.rgb = MUTED
        p_cap.alignment = PP_ALIGN.CENTER

    # =========================================================================
    # SLIDE 5: TECHNOLOGY AND IMPLEMENTATION
    # =========================================================================
    slide5 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide5, "Technology and implementation", "Full Stack Specifications")

    # Table of Tech Stack
    rows, cols = 6, 3
    left, top, width, height = Inches(0.8), Inches(1.6), Inches(11.733), Inches(3.4)
    table_shape = slide5.shapes.add_table(rows, cols, left, top, width, height)
    table = table_shape.table
    table.columns[0].width = Inches(2.2)
    table.columns[1].width = Inches(3.6)
    table.columns[2].width = Inches(5.933)

    tech_headers = ["LAYER", "SAMPLE CHOICE", "PURPOSE"]
    tech_data = [
        ["Frontend (BFF)", "Next.js 15, React 19, Tailwind CSS", "Server-Side Rendering (SSR), App Router, responsive storefront, and HttpOnly session cookies."],
        ["API Gateway Edge", "Node.js, Express, Helmet, AJV", "Centralized routing, sliding-window rate limiting (120 req/60s), schema guard, and circuit breaking."],
        ["Domain Microservices", "Express.js REST & GraphQL", "7 autonomous microservices (Auth, Catalog, Cart, Wishlist, Order Saga, Payment, Notification)."],
        ["Persistence Stores", "MongoDB Community 7.0 & MySQL 8.4", "Zero-shared databases: MongoDB collections per service + MySQL for durable edge idempotency."],
        ["Partner Integrations", "Razorpay REST API & Resend Mailer", "UPI/Card digital checkouts in Indian Paise, HMAC-SHA256 signature verification, and transactional emails."]
    ]

    for c_idx, h_text in enumerate(tech_headers):
        cell = table.cell(0, c_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = ACCENT_BLUE
        p = cell.text_frame.paragraphs[0]
        p.text = h_text
        p.font.name = "Calibri"
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = WHITE
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE

    for r_idx, row_values in enumerate(tech_data):
        for c_idx, val in enumerate(row_values):
            cell = table.cell(r_idx + 1, c_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = LIGHT_BG if r_idx % 2 == 1 else WHITE
            p = cell.text_frame.paragraphs[0]
            p.text = val
            p.font.name = "Calibri"
            p.font.size = Pt(10.5)
            p.font.color.rgb = NAVY if c_idx == 0 else DARK
            p.font.bold = (c_idx == 0)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE

    # Technical Highlight Card
    card_tech = create_card(slide5, Inches(0.8), Inches(5.2), Inches(11.733), Inches(1.7), bg_color=WHITE, border_color=BLUE)
    tb_tp = slide5.shapes.add_textbox(Inches(1.0), Inches(5.3), Inches(11.3), Inches(1.5))
    tf_tp = tb_tp.text_frame
    tf_tp.word_wrap = True

    p_tph = tf_tp.paragraphs[0]
    p_tph.text = "One technical point to discuss: Durable Idempotency & Distributed Rollback"
    p_tph.font.name = "Calibri"
    p_tph.font.size = Pt(12)
    p_tph.font.bold = True
    p_tph.font.color.rgb = BLUE
    p_tph.space_after = Pt(4)

    p_tp1 = tf_tp.add_paragraph()
    p_tp1.text = "1. Duplicate Order Prevention: Clients pass an Idempotency-Key header on checkout. The API Gateway computes a SHA-256 hash of the payload and queries the MySQL api_idempotency table. Repeated requests return cached responses with Idempotency-Replayed: true, completely preventing double-billing on network retries."
    p_tp1.font.name = "Calibri"
    p_tp1.font.size = Pt(10.5)
    p_tp1.font.color.rgb = DARK
    p_tp1.space_after = Pt(3)

    p_tp2 = tf_tp.add_paragraph()
    p_tp2.text = "2. Compensatory Stock Rollback: During order placement, Product Service inventory is locked atomically. If Razorpay payment fails or signature validation mismatches, the Order Service executes a compensatory REST call to release the reserved stock immediately."
    p_tp2.font.name = "Calibri"
    p_tp2.font.size = Pt(10.5)
    p_tp2.font.color.rgb = DARK

    # =========================================================================
    # SLIDE 6: SCREENSHOTS AND DEMONSTRATION
    # =========================================================================
    slide6 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide6, "Screenshots and demonstration", "Operational Application Walkthrough")

    # Subtitle
    tb_sub6 = slide6.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(11.7), Inches(0.35))
    p_s6 = tb_sub6.text_frame.paragraphs[0]
    p_s6.text = "Live validation of core storefront and administrative fulfillment workflows"
    p_s6.font.name = "Calibri"
    p_s6.font.size = Pt(12.5)
    p_s6.font.color.rgb = DARK

    # Left Card: Screen 1
    create_card(slide6, Inches(0.8), Inches(1.85), Inches(5.6), Inches(4.3))
    tb_sc1 = slide6.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(5.2), Inches(4.0))
    tf_sc1 = tb_sc1.text_frame
    tf_sc1.word_wrap = True
    
    p_sc1_h = tf_sc1.paragraphs[0]
    p_sc1_h.text = "01  Storefront Catalog & Razorpay Checkout"
    p_sc1_h.font.name = "Calibri"
    p_sc1_h.font.size = Pt(13)
    p_sc1_h.font.bold = True
    p_sc1_h.font.color.rgb = NAVY
    p_sc1_h.space_after = Pt(8)

    # Inner preview box
    sc1_box = slide6.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(2.55), Inches(5.2), Inches(2.3))
    sc1_box.fill.solid()
    sc1_box.fill.fore_color.rgb = WHITE
    sc1_box.line.color.rgb = CARD_BORDER
    tb_in1 = slide6.shapes.add_textbox(Inches(1.1), Inches(2.7), Inches(5.0), Inches(2.0))
    tf_in1 = tb_in1.text_frame
    tf_in1.word_wrap = True
    p_in1 = tf_in1.paragraphs[0]
    p_in1.text = "[ Storefront UI / Next.js 15 Web Application ]\n\n• Dynamic product grid with rating badges & stock status\n• Slide-over interactive cart drawer with live subtotal\n• Razorpay Standard Checkout popup supporting UPI & Cards\n• Server-validated coupon promotional discounts"
    p_in1.font.name = "Calibri"
    p_in1.font.size = Pt(10)
    p_in1.font.color.rgb = DARK

    p_sc1_d = tf_sc1.add_paragraph()
    p_sc1_d.text = "\nDemonstrates instant catalog filtering, dynamic ETag validation (304), persistent carts, and Razorpay modal checkout."
    p_sc1_d.font.name = "Calibri"
    p_sc1_d.font.size = Pt(10.5)
    p_sc1_d.font.color.rgb = MUTED

    # Right Card: Screen 2
    create_card(slide6, Inches(6.9), Inches(1.85), Inches(5.6), Inches(4.3))
    tb_sc2 = slide6.shapes.add_textbox(Inches(7.1), Inches(2.0), Inches(5.2), Inches(4.0))
    tf_sc2 = tb_sc2.text_frame
    tf_sc2.word_wrap = True
    
    p_sc2_h = tf_sc2.paragraphs[0]
    p_sc2_h.text = "02  Order Tracking & Staff Inventory Workbench"
    p_sc2_h.font.name = "Calibri"
    p_sc2_h.font.size = Pt(13)
    p_sc2_h.font.bold = True
    p_sc2_h.font.color.rgb = NAVY
    p_sc2_h.space_after = Pt(8)

    # Inner preview box
    sc2_box = slide6.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(7.1), Inches(2.55), Inches(5.2), Inches(2.3))
    sc2_box.fill.solid()
    sc2_box.fill.fore_color.rgb = WHITE
    sc2_box.line.color.rgb = CARD_BORDER
    tb_in2 = slide6.shapes.add_textbox(Inches(7.2), Inches(2.7), Inches(5.0), Inches(2.0))
    tf_in2 = tb_in2.text_frame
    tf_in2.word_wrap = True
    p_in2 = tf_in2.paragraphs[0]
    p_in2.text = "[ Customer Tracking & Admin Workbench ]\n\n• Delivery timeline stepper (Confirmed → Shipped → Delivered)\n• Staff stock management table with mandatory reason prompt\n• Automated return policy verification (30-day window check)\n• Refund authorization dashboard triggering Razorpay API"
    p_in2.font.name = "Calibri"
    p_in2.font.size = Pt(10)
    p_in2.font.color.rgb = DARK

    p_sc2_d = tf_sc2.add_paragraph()
    p_sc2_d.text = "\nDemonstrates immutable inventory stockHistory[], courier tracking updates, and 30-day return claims."
    p_sc2_d.font.name = "Calibri"
    p_sc2_d.font.size = Pt(10.5)
    p_sc2_d.font.color.rgb = MUTED

    # Bottom Demo Sequence Banner
    demo_card = create_card(slide6, Inches(0.8), Inches(6.3), Inches(11.7), Inches(0.75), bg_color=WHITE, border_color=ACCENT_BLUE)
    tb_demo = slide6.shapes.add_textbox(Inches(1.0), Inches(6.35), Inches(11.3), Inches(0.6))
    p_dm = tb_demo.text_frame.paragraphs[0]
    p_dm.text = "Demo sequence:  "
    p_dm.font.name = "Calibri"
    p_dm.font.size = Pt(11)
    p_dm.font.bold = True
    p_dm.font.color.rgb = BLUE
    
    r_dm = p_dm.add_run()
    r_dm.text = "Login & Authenticate   →   Search Products   →   Add to Cart   →   Razorpay Payment (UPI)   →   Verify Order Tracking"
    r_dm.font.bold = True
    r_dm.font.size = Pt(11)
    r_dm.font.color.rgb = NAVY

    # =========================================================================
    # SLIDE 7: TESTING AND RESULTS
    # =========================================================================
    slide7 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide7, "Testing and results", "Verification Evidence")

    # Subtitle
    tb_sub7 = slide7.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(11.7), Inches(0.35))
    p_s7 = tb_sub7.text_frame.paragraphs[0]
    p_s7.text = "Verification evidence that core distributed microservice functions work reliably"
    p_s7.font.name = "Calibri"
    p_s7.font.size = Pt(12.5)
    p_s7.font.color.rgb = DARK

    # Table of Test Results
    t_rows, t_cols = 4, 3
    t_left, t_top, t_w, t_h = Inches(0.8), Inches(1.85), Inches(11.733), Inches(4.8)
    table_test = slide7.shapes.add_table(t_rows, t_cols, t_left, t_top, t_w, t_h).table
    table_test.columns[0].width = Inches(2.8)
    table_test.columns[1].width = Inches(4.8)
    table_test.columns[2].width = Inches(4.133)

    test_headers = ["TEST CASE", "EXPECTED BEHAVIOR", "OBSERVED RESULT"]
    test_data = [
        [
            "User Authentication & RBAC Session Security",
            "Bcrypt credentials checked; issue JWT strictly in HttpOnly SameSite=Lax cookie. Reject unauthenticated/unauthorized access with RFC 7807 problem details (HTTP 401 / 403).",
            "PASSED (100%)\nToken successfully set in HttpOnly cookie; unprivileged role access returns HTTP 403 Forbidden with uniform problem details payload."
        ],
        [
            "Checkout Saga & Stock Rollback on Payment Failure",
            "Atomically reserve warehouse stock during order placement; if Razorpay payment is declined or aborted, trigger compensatory REST call to release reserved stock back.",
            "PASSED (100%)\nStock decremented upon checkout; intentional payment failure successfully invoked compensatory rollback, restoring warehouse stock count."
        ],
        [
            "Razorpay Payment Signature Verification (HMAC)",
            "Compute HMAC-SHA256(order_id + '|' + payment_id, SECRET). Match signature string to confirm payment. Tampered signatures must immediately set payment status to FAILED.",
            "PASSED (100%)\nValid signatures transition order to CONFIRMED. Altered signatures reject with HTTP 400 Bad Request (urn:smartcart:error:payment-declined)."
        ]
    ]

    for c_idx, h_text in enumerate(test_headers):
        cell = table_test.cell(0, c_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = ACCENT_BLUE
        p = cell.text_frame.paragraphs[0]
        p.text = h_text
        p.font.name = "Calibri"
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = WHITE
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE

    for r_idx, row_values in enumerate(test_data):
        for c_idx, val in enumerate(row_values):
            cell = table_test.cell(r_idx + 1, c_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = LIGHT_BG if r_idx % 2 == 1 else WHITE
            p = cell.text_frame.paragraphs[0]
            p.text = val
            p.font.name = "Calibri"
            p.font.size = Pt(10)
            if c_idx == 0:
                p.font.bold = True
                p.font.color.rgb = NAVY
            elif c_idx == 2:
                p.font.bold = True
                p.font.color.rgb = GREEN
            else:
                p.font.color.rgb = DARK
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE

    # =========================================================================
    # SLIDE 8: CONCLUSION AND NEXT STEPS
    # =========================================================================
    slide8 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide8, "Conclusion and next steps", "Milestones & Future Scope")

    # Left Column: What the team completed
    card_comp = create_card(slide8, Inches(0.8), Inches(1.7), Inches(5.6), Inches(5.1))
    tb_comp = slide8.shapes.add_textbox(Inches(1.1), Inches(1.9), Inches(5.0), Inches(4.7))
    tf_comp = tb_comp.text_frame
    tf_comp.word_wrap = True
    
    p_ch = tf_comp.paragraphs[0]
    p_ch.text = "WHAT THE TEAM COMPLETED"
    p_ch.font.name = "Calibri"
    p_ch.font.size = Pt(14)
    p_ch.font.bold = True
    p_ch.font.color.rgb = BLUE
    p_ch.space_after = Pt(14)

    comp_items = [
        ("Autonomous Microservices Architecture:", "Built and deployed 8 standalone services (Auth, Catalog, Cart, Wishlist, Order Saga, Payment, Notification, API Gateway) with isolated database schemas."),
        ("End-to-End E-Commerce Workflows:", "Full product catalog with dynamic ETag caching (304), server-authoritative cart calculation, coupon promotion engine, and 30-day return workflows."),
        ("Production Payment & Edge Hardening:", "Completed Razorpay standard checkout integration in Paise, HMAC-SHA256 signature verification, durable MySQL idempotency, rate limiting, and RFC 7807 errors.")
    ]
    for title, desc in comp_items:
        p_t = tf_comp.add_paragraph()
        p_t.text = "• " + title + " "
        p_t.font.name = "Calibri"
        p_t.font.bold = True
        p_t.font.size = Pt(11.5)
        p_t.font.color.rgb = NAVY
        
        run_d = p_t.add_run()
        run_d.text = desc
        run_d.font.bold = False
        run_d.font.size = Pt(11)
        run_d.font.color.rgb = DARK
        p_t.space_after = Pt(12)

    # Right Column: What comes next
    card_next = create_card(slide8, Inches(6.9), Inches(1.7), Inches(5.6), Inches(5.1))
    tb_next = slide8.shapes.add_textbox(Inches(7.2), Inches(1.9), Inches(5.0), Inches(4.7))
    tf_next = tb_next.text_frame
    tf_next.word_wrap = True
    
    p_nh = tf_next.paragraphs[0]
    p_nh.text = "WHAT COMES NEXT"
    p_nh.font.name = "Calibri"
    p_nh.font.size = Pt(14)
    p_nh.font.bold = True
    p_nh.font.color.rgb = ACCENT_BLUE
    p_nh.space_after = Pt(14)

    next_items = [
        ("Asynchronous Message Broker Integration:", "Transition cross-service communication from synchronous HTTP/REST to an event-driven pub/sub architecture using RabbitMQ or Apache Kafka."),
        ("Multi-Channel Logistics Webhooks:", "Integrate third-party carrier APIs (Shiprocket / Bluedart) for automated live GPS parcel tracking webhooks instead of manual status advancement."),
        ("Multi-Channel SMS Alerts & Search AI:", "Add SMS dispatch via Twilio alongside email alerts, and integrate vector-based AI search for natural language product discovery.")
    ]
    for title, desc in next_items:
        p_t = tf_next.add_paragraph()
        p_t.text = "• " + title + " "
        p_t.font.name = "Calibri"
        p_t.font.bold = True
        p_t.font.size = Pt(11.5)
        p_t.font.color.rgb = NAVY
        
        run_d = p_t.add_run()
        run_d.text = desc
        run_d.font.bold = False
        run_d.font.size = Pt(11)
        run_d.font.color.rgb = DARK
        p_t.space_after = Pt(12)

    # =========================================================================
    # SLIDE 9: THANK YOU
    # =========================================================================
    slide9 = prs.slides.add_slide(blank_slide_layout)
    
    bg9 = slide9.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg9.fill.solid()
    bg9.fill.fore_color.rgb = WHITE
    bg9.line.fill.background()

    tb_ty = slide9.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.333), Inches(1.2))
    p_ty = tb_ty.text_frame.paragraphs[0]
    p_ty.text = "Thank you"
    p_ty.font.name = "Calibri"
    p_ty.font.size = Pt(44)
    p_ty.font.bold = True
    p_ty.font.color.rgb = NAVY
    p_ty.alignment = PP_ALIGN.CENTER

    tb_q = slide9.shapes.add_textbox(Inches(1.0), Inches(3.0), Inches(11.333), Inches(0.5))
    p_q = tb_q.text_frame.paragraphs[0]
    p_q.text = "Questions & Answers / Technical Discussion"
    p_q.font.name = "Calibri"
    p_q.font.size = Pt(18)
    p_q.font.color.rgb = BLUE
    p_q.alignment = PP_ALIGN.CENTER

    # Project Links Card
    card_links = create_card(slide9, Inches(3.0), Inches(3.8), Inches(7.333), Inches(2.2), bg_color=LIGHT_BG)
    tb_l = slide9.shapes.add_textbox(Inches(3.3), Inches(4.0), Inches(6.733), Inches(1.8))
    tf_l = tb_l.text_frame
    tf_l.word_wrap = True

    pl1 = tf_l.paragraphs[0]
    pl1.text = "SmartCart Project Deliverables & Repository:"
    pl1.font.bold = True
    pl1.font.size = Pt(12)
    pl1.font.color.rgb = NAVY
    pl1.space_after = Pt(8)

    pl2 = tf_l.add_paragraph()
    pl2.text = "• GitHub: https://github.com/Ayush17314/SmartCart"
    pl2.font.size = Pt(11.5)
    pl2.font.color.rgb = DARK
    pl2.space_after = Pt(4)

    pl3 = tf_l.add_paragraph()
    pl3.text = "• Engineering Team: Group 03 (TeamID 03) — Ayush Dubey"
    pl3.font.size = Pt(11.5)
    pl3.font.color.rgb = DARK
    pl3.space_after = Pt(4)

    pl4 = tf_l.add_paragraph()
    pl4.text = "• IIIT Vadodara  |  CS 545 Software Engineering  |  Autumn 2026-27"
    pl4.font.bold = True
    pl4.font.size = Pt(11)
    pl4.font.color.rgb = MUTED

    # Save presentation
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    prs.save(output_path)
    print(f"Presentation successfully saved at: {output_path}")

if __name__ == "__main__":
    out = os.path.join("docs", "SmartCart_Project_Presentation.pptx")
    create_smartcart_presentation(out)
